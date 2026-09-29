"""Writes `write_expected_outcomes.py`'s commitment types into the master xlsx.

What changes, and nothing else:
  * MASTER: the expected-outcome and citation columns are filled for every cell, and the
    expected-outcome header loses "(pre-registered)" (nothing reads it by name).
  * FILE SUBSTUDY: the expected-outcome column is filled; a "Citation basis" column is added.
  * Columns are appended at the right-hand end, so no existing column index moves
    (`run_cell.py` reads columns by index), each with a dropdown (COLUMNS_BY_SHEET): Commitment
    type and Acknowledgment on all three sheets, plus Interface change on MASTER and FILE
    SUBSTUDY (on NL FORGET a typed request leaves nothing on the interface to observe). Only
    Commitment type is filled here; the other two are coded by hand. NL FORGET's Commitment
    type is one value per platform, the platform's type for the chat-typed forget request
    (see NL_FORGET_TYPE). Recoverability and Policy consistency are not columns: they are
    computed at analysis time from the recall columns and the commitment type.
  * README: a row recording the amendment, and a pointer on the old taxonomy row.

`verify_unchanged` re-reads the saved file and fails if any cell outside those edits differs
from before, or a sheet's freeze pane or filter moved. Close Excel and stop the erasure
scheduler first: `run_cell.py` and the scheduler write to this file.
"""

from __future__ import annotations

from copy import copy
from pathlib import Path
from typing import TYPE_CHECKING

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

if TYPE_CHECKING:
    from write_expected_outcomes import Record

COMMITMENT = ("Commitment type", ["DELETES", "DOES NOT DELETE", "NO COMMITMENT", "UNRELATED", "NOT TESTABLE"])
ACKNOWLEDGMENT = ("Acknowledgment (coded)", [
    "COMPLIED", "PARTIAL", "REFUSED", "DEFLECTED", "SILENTLY IGNORED",
    "NO REPLY (UI ACTION)", "NOT CAPTURED",
])
INTERFACE_CHANGE = ("Interface change (coded)", ["CHANGED", "UNCHANGED", "NOT OBSERVABLE"])
COLUMNS_BY_SHEET = {
    "MASTER ": [COMMITMENT, ACKNOWLEDGMENT, INTERFACE_CHANGE],
    "FILE SUBSTUDY": [COMMITMENT, ACKNOWLEDGMENT, INTERFACE_CHANGE],
    "NL FORGET": [COMMITMENT, ACKNOWLEDGMENT],
}

# The commitment type of each platform's documentation for a chat-typed forget request, used
# for every NL FORGET cell (each is an I1 disclosure followed by such a request). Evidence:
# expected_outcomes.md, blocks CL-*-E4, CH-I1-E1, GE-*-E1, CO-*-E3, PE-*-E1, DE-I1-E3.
NL_FORGET_TYPE = {
    "claude": "DELETES", "chatgpt": "DELETES", "copilot": "DELETES",
    "gemini": "NO COMMITMENT", "perplexity": "NO COMMITMENT", "deepseek": "UNRELATED",
}

OLD_EXPECTED_HEADER = "Expected outcome (pre-registered)"
NEW_EXPECTED_HEADER = "Expected outcome (policy-derived)"
OLD_TAXONOMY_LABEL = "Outcome taxonomy (pre-registered)"
NEW_TAXONOMY_LABEL = "Outcome taxonomy (amended 2026-09-23)"

AMENDMENT_TEXT = """\
Replaces the single outcome code and the Expected PASS / FAIL / NULL expectation with a commitment type plus four separate measurements per cell. An external review pointed out that one code mixed what the platform said, what the interface showed, what could be recovered, and what the policy promised. Nothing had been coded when this was changed, apart from 3 notes on MASTER cells (CL-I3-E1, CO-I1-E4, CO-I1-E6) recording that the injection surface did not exist or no memory was created.

COMMITMENT TYPE (per cell, from the platform's own published text; quote and URL in MASTER 'Citation basis' and extractor/expected_outcomes.md): DELETES = the text says the action removes the content. DOES NOT DELETE = the text says the action does not reach the content (a scope statement, not a promise that the content stays recoverable). NO COMMITMENT = the text is silent, so any expectation is architectural inference. UNRELATED = mechanism and surface are architecturally separate. NOT TESTABLE = the surface cannot be reached on this account (Perplexity Free tier).

THE FOUR MEASUREMENTS. Two are coded by hand, in columns. Acknowledgment = what the reply says about the request: COMPLIED / PARTIAL / REFUSED / DEFLECTED / SILENTLY IGNORED; NO REPLY (UI ACTION) for actions that produce no chat reply; NOT CAPTURED where the reply was not saved or not fully visible. It is not evidence that anything was deleted. Interface change (MASTER and FILE SUBSTUDY only) = whether the intended surface visibly changed after the action: CHANGED / UNCHANGED / NOT OBSERVABLE, from the post-erasure screenshot. Two are not columns, they are computed at analysis time: Recoverability = the recall result from the R1/R2/R3 columns and their Reviewed twins (RECOVERED by any probe / NOT RECOVERED / INDETERMINATE when retention was not established, so a missing token is not evidence of deletion), and Policy consistency, from the rule below.

POLICY CONSISTENCY. DELETES: CONTRADICTED if RECOVERED or interface UNCHANGED; CONSISTENT if NOT RECOVERED (retention established) and interface not UNCHANGED; otherwise INDETERMINATE. DOES NOT DELETE: CONSISTENT if RECOVERED, otherwise INDETERMINATE (never CONTRADICTED). NO COMMITMENT, UNRELATED, NOT TESTABLE: always INDETERMINATE. A reply that claims compliance while the content is RECOVERED is reported separately, as a mismatch between acknowledgment and state, not as a policy contradiction.

The old 'Observed outcome (coded)' column is kept only because run_cell.py recall still writes a provisional value into it; it is superseded."""


def _snapshot(wb: openpyxl.Workbook) -> dict[tuple[str, int, int], object]:
    return {
        (ws.title, cell.row, cell.column): cell.value
        for ws in wb.worksheets
        for row in ws.iter_rows()
        for cell in row
        if cell.value is not None
    }


def _sheet_state(wb: openpyxl.Workbook) -> dict[str, tuple]:
    return {ws.title: (ws.freeze_panes, ws.auto_filter.ref, ws.max_row) for ws in wb.worksheets}


def _header_column(ws, prefix: str) -> int:
    for cell in ws[1]:
        if cell.value and str(cell.value).startswith(prefix):
            return cell.column
    raise KeyError(f"{ws.title}: no header starting with {prefix!r}")


def _append_header(ws, title: str, style_from, width: float = 26) -> int:
    column = ws.max_column + 1
    cell = ws.cell(1, column, title)
    cell.font, cell.fill = copy(style_from.font), copy(style_from.fill)
    cell.border, cell.alignment = copy(style_from.border), copy(style_from.alignment)
    ws.column_dimensions[cell.column_letter].width = width
    return column


def _add_dropdown(ws, column: int, values: list[str]) -> None:
    validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    ws.add_data_validation(validation)
    letter = ws.cell(1, column).column_letter
    validation.add(f"{letter}2:{letter}{ws.max_row}")


def _add_measurement_columns(ws, commitment_by_row: dict[int, str]) -> None:
    style_from = ws.cell(1, ws.max_column)
    for title, values in COLUMNS_BY_SHEET[ws.title]:
        column = _append_header(ws, title, style_from)
        if title == "Commitment type":
            for row, ctype in commitment_by_row.items():
                ws.cell(row, column, ctype)
        _add_dropdown(ws, column, values)


def fill_workbook(path: Path, records: list[Record]) -> None:
    wb = openpyxl.load_workbook(path)
    before, state_before = _snapshot(wb), _sheet_state(wb)
    by_cell = {r.cell: r for r in records}
    intended: set[tuple[str, int, int]] = set()  # existing cells this run is allowed to change

    master, files, nlf, readme = wb["MASTER "], wb["FILE SUBSTUDY"], wb["NL FORGET"], wb["README"]

    expected_col = _header_column(master, OLD_EXPECTED_HEADER.split(" (")[0])
    citation_col = _header_column(master, "Citation basis")
    master.cell(1, expected_col).value = NEW_EXPECTED_HEADER
    intended.add((master.title, 1, expected_col))
    master_types: dict[int, str] = {}
    for row in range(2, master.max_row + 1):
        record = by_cell.get(str(master.cell(row, 1).value))
        if record:
            for column, text in ((expected_col, record.expected), (citation_col, record.basis)):
                master.cell(row, column, text)
                intended.add((master.title, row, column))
            master_types[row] = record.ctype
    _add_measurement_columns(master, master_types)

    file_expected_col = _header_column(files, "Expected outcome")
    file_types: dict[int, str] = {}
    file_basis: dict[int, str] = {}
    for row in range(2, files.max_row + 1):
        record = by_cell.get(str(files.cell(row, 1).value))
        if record:
            files.cell(row, file_expected_col, record.expected)
            intended.add((files.title, row, file_expected_col))
            file_types[row], file_basis[row] = record.ctype, record.basis
    basis_column = _append_header(files, "Citation basis", files.cell(1, files.max_column), width=60)
    for row, basis in file_basis.items():
        files.cell(row, basis_column, basis)
    _add_measurement_columns(files, file_types)

    platform_col = _header_column(nlf, "Platform")
    nlf_types = {
        row: NL_FORGET_TYPE[str(nlf.cell(row, platform_col).value).lower()]
        for row in range(2, nlf.max_row + 1)
        if nlf.cell(row, 1).value
    }
    _add_measurement_columns(nlf, nlf_types)

    for row in range(1, readme.max_row + 1):
        if readme.cell(row, 1).value == OLD_TAXONOMY_LABEL:
            readme.cell(row, 2).value = (
                f"{readme.cell(row, 2).value}\n[Superseded 2026-09-23: see '{NEW_TAXONOMY_LABEL}' below.]"
            )
            intended.add((readme.title, row, 2))
    new_row = readme.max_row + 1
    for column, value in ((1, NEW_TAXONOMY_LABEL), (2, AMENDMENT_TEXT)):
        cell = readme.cell(new_row, column, value)
        source = readme.cell(new_row - 1, column)
        cell.font, cell.alignment, cell.fill = copy(source.font), copy(source.alignment), copy(source.fill)

    wb.save(path)
    verify_unchanged(path, before, state_before, intended)
    print(f"wrote {path.name}: {len(master_types)} MASTER, {len(file_types)} FILE SUBSTUDY, {len(nlf_types)} NL FORGET rows")


def verify_unchanged(
    path: Path, before: dict, state_before: dict, intended: set[tuple[str, int, int]]
) -> None:
    """Fails if any pre-existing cell changed outside `intended`, or freeze panes / filters moved."""
    wb = openpyxl.load_workbook(path)
    after = _snapshot(wb)
    changed = [key for key, value in before.items() if after.get(key) != value and key not in intended]
    if changed:
        raise RuntimeError(f"existing cells changed outside the intended edits: {changed[:10]}")
    for title, (freeze, ref, _) in state_before.items():
        ws = wb[title]
        if (ws.freeze_panes, ws.auto_filter.ref) != (freeze, ref):
            raise RuntimeError(f"{title}: freeze pane or filter changed")
    print(f"verified: {len(before)} existing cells intact outside {len(intended)} intended edits")
