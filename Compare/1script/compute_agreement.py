"""Compute inter-rater agreement (raw % and Gwet's AC1) between
Taiwo_Coding.xlsx and Jihwan_Coding.xlsx for each codebook dimension, and
regenerate agreement_report.md's summary table + a machine-readable CSV.

Usage: python3 compute_agreement.py
"""

import csv
from collections import Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
TAIWO_FILE = HERE / "Taiwo_Coding.xlsx"
JIHWAN_FILE = HERE / "Jihwan_Coding.xlsx"
REPORT_FILE = HERE / "agreement_report.md"
CSV_OUT = HERE / "agreement_results.csv"

DIMENSIONS = [
    "deletion_location",
    "justification",
    "what",
    "mood",
    "tone",
    "verb",
    "accompanying_request",
]

SHEET_NAME = "Coding"
N_ROWS = 50  # first-pass shared n=50 comparison set


def load_rows(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET_NAME]
    header = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=1 + N_ROWS, values_only=True):
        rows.append(dict(zip(header, row)))
    return rows


def normalize(value) -> str:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip().upper()


def gwet_ac1(t_vals: list[str], j_vals: list[str], categories: list[str], n: int) -> tuple[float, float]:
    c1, c2 = Counter(t_vals), Counter(j_vals)
    po = sum(1 for t, j in zip(t_vals, j_vals) if t == j) / n
    q = len(categories)
    pbar = {c: ((c1[c] / n) + (c2[c] / n)) / 2 for c in categories}
    pe_gwet = sum(pbar[c] * (1 - pbar[c]) for c in categories) / (q - 1)
    ac1 = (po - pe_gwet) / (1 - pe_gwet)
    return po, ac1


def compute_dimension(taiwo: list[dict], jihwan: list[dict], dim: str) -> dict:
    t_vals = [normalize(r[dim]) for r in taiwo]
    j_vals = [normalize(r[dim]) for r in jihwan]
    n = len(t_vals)
    categories = sorted(set(t_vals + j_vals))

    agree_count = sum(1 for t, j in zip(t_vals, j_vals) if t == j)
    po, ac1 = gwet_ac1(t_vals, j_vals, categories, n)

    return {
        "dimension": dim,
        "n": n,
        "agree_count": agree_count,
        "pct_agreement": round(100 * po, 1),
        "ac1": round(ac1, 3),
    }


def render_table(results: list[dict]) -> str:
    lines = [
        "| Dimension | Agreement | % | Gwet's AC1 |",
        "|---|---|---|---|",
    ]
    for r in results:
        lines.append(
            f"| {r['dimension']} | {r['agree_count']}/{r['n']} | "
            f"{r['pct_agreement']}% | {r['ac1']} |"
        )
    return "\n".join(lines)


def update_report(results: list[dict]) -> None:
    text = REPORT_FILE.read_text()
    start = text.index("| Dimension |")
    end = text.index("\n\n", start)
    new_table = render_table(results)
    REPORT_FILE.write_text(text[:start] + new_table + text[end:])


def write_csv(results: list[dict]) -> None:
    with CSV_OUT.open("w", newline="") as f:
        writer = csv.DictWriter(
            f, fieldnames=["dimension", "n", "agree_count", "pct_agreement", "ac1"]
        )
        writer.writeheader()
        writer.writerows(results)


def compute_overall_micro(taiwo: list[dict], jihwan: list[dict]) -> dict:
    """Pool all (dimension, item) rating pairs into one set, the same way the
    existing micro-agreement % is computed (306/350), and report AC1 on that
    pooled set. Category labels are prefixed with their dimension so e.g.
    verb's 'none' and accompanying_request's 'none' aren't conflated."""
    t_vals: list[str] = []
    j_vals: list[str] = []
    for dim in DIMENSIONS:
        t_vals += [f"{dim}::{normalize(r[dim])}" for r in taiwo]
        j_vals += [f"{dim}::{normalize(r[dim])}" for r in jihwan]

    n = len(t_vals)
    categories = sorted(set(t_vals + j_vals))
    agree_count = sum(1 for t, j in zip(t_vals, j_vals) if t == j)
    po, ac1 = gwet_ac1(t_vals, j_vals, categories, n)

    return {
        "n": n,
        "agree_count": agree_count,
        "pct_agreement": round(100 * po, 1),
        "ac1": round(ac1, 3),
    }


def update_overall_line(overall: dict) -> None:
    text = REPORT_FILE.read_text()
    start = text.index("**Overall micro agreement:**")
    end = text.index("\n", start)
    new_line = (
        f"**Overall micro agreement:** {overall['agree_count']}/{overall['n']} "
        f"= {overall['pct_agreement']}% "
        f"(pooled Gwet's AC1 = {overall['ac1']})"
    )
    text = text[:start] + new_line + text[end:]
    REPORT_FILE.write_text(text)


def main() -> None:
    taiwo = load_rows(TAIWO_FILE)
    jihwan = load_rows(JIHWAN_FILE)

    results = [compute_dimension(taiwo, jihwan, dim) for dim in DIMENSIONS]
    overall = compute_overall_micro(taiwo, jihwan)

    print(render_table(results))
    print(
        f"\nOverall micro: {overall['agree_count']}/{overall['n']} = "
        f"{overall['pct_agreement']}% | pooled AC1 = {overall['ac1']}"
    )

    update_report(results)
    update_overall_line(overall)
    write_csv(results)
    print(f"\nUpdated {REPORT_FILE.name} and wrote {CSV_OUT.name}")


if __name__ == "__main__":
    main()
