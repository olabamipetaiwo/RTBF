"""Fourth Pass inter-rater agreement: all 138 prompts, all 7 dimensions.

Gwet's AC1 only (professor directive 2026-09-11); Cohen's kappa is
intentionally not computed.

Sources: ./Taiwo_Coding.xlsx and ./Jihwan_Coding_Final.xlsx (sheet "Coding").
Method matches Third Pass: categories are the union of both raters' labels per
dimension; the pooled micro figure prefixes each label with its dimension so
equal labels in different dimensions (e.g. "none") are not conflated.

Writes agreement_report.md and disagreements.csv here.

Usage: python3 compute_agreement.py
"""

import csv
from collections import Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
TAIWO_FILE = HERE / "Taiwo_Coding.xlsx"
JIHWAN_FILE = HERE / "Jihwan_Coding_Final.xlsx"
SHEET = "Coding"
DIMENSIONS = [
    "deletion_location",
    "justification",
    "what",
    "mood",
    "tone",
    "verb",
    "accompanying_request",
]


def normalize(value: object) -> str:
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    return str(value).strip().upper()


def load(path: Path) -> dict[int, dict[str, object]]:
    ws = openpyxl.load_workbook(path)[SHEET]
    header = [c.value for c in ws[1]]
    return {
        int(row[0]): dict(zip(header, row))
        for row in ws.iter_rows(min_row=2, values_only=True)
        if row[0] is not None
    }


def observed_and_marginals(a: list[str], b: list[str]) -> tuple[float, dict[str, float]]:
    n = len(a)
    ca, cb = Counter(a), Counter(b)
    po = sum(x == y for x, y in zip(a, b)) / n
    pi = {c: (ca[c] + cb[c]) / (2 * n) for c in set(a) | set(b)}
    return po, pi


def gwet_ac1(a: list[str], b: list[str]) -> float:
    po, pi = observed_and_marginals(a, b)
    if len(pi) < 2:
        return 1.0
    pe = sum(p * (1 - p) for p in pi.values()) / (len(pi) - 1)
    return (po - pe) / (1 - pe)


def main() -> None:
    taiwo, jihwan = load(TAIWO_FILE), load(JIHWAN_FILE)
    assert taiwo.keys() == jihwan.keys(), "item_no sets differ between books"
    items = sorted(taiwo)
    n = len(items)

    stats, dis_rows = [], []
    pooled_t: list[str] = []
    pooled_j: list[str] = []
    for dim in DIMENSIONS:
        a = [normalize(taiwo[i][dim]) for i in items]
        b = [normalize(jihwan[i][dim]) for i in items]
        agree = sum(x == y for x, y in zip(a, b))
        stats.append((dim, agree, gwet_ac1(a, b)))
        pooled_t += [f"{dim}::{x}" for x in a]
        pooled_j += [f"{dim}::{x}" for x in b]
        dis_rows += [
            [dim, i, str(taiwo[i]["prompt"]), x.lower(), y.lower()]
            for i, x, y in zip(items, a, b)
            if x != y
        ]

    tot = len(pooled_t)
    p_agree = sum(x == y for x, y in zip(pooled_t, pooled_j))
    p_ac1 = gwet_ac1(pooled_t, pooled_j)

    with (HERE / "disagreements.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["dimension", "item_no", "prompt", "taiwo", "jihwan", "adjudicated_value", "adjudication_notes"])
        w.writerows([*row, "", ""] for row in dis_rows)

    lines = [
        f"# Final Inter-Rater Agreement, Fourth Pass (Gwet's AC1, n={n} per dimension)",
        "",
        "| Dimension | Agreement | % | Gwet's AC1 |",
        "|---|---|---|---|",
    ]
    lines += [f"| {d} | {ag}/{n} | {100 * ag / n:.1f}% | {ac1:.3f} |" for d, ag, ac1 in stats]
    lines += [
        "",
        f"**Overall micro agreement:** {p_agree}/{tot} = {100 * p_agree / tot:.1f}% ",
        f"**Gwet's AC1:** {p_ac1:.3f}",
        "",
        "## Disagreements by dimension",
        "",
    ]
    for dim in DIMENSIONS:
        rows = [r for r in dis_rows if r[0] == dim]
        lines += [
            f"### {dim} ({len(rows)} disagreements)",
            "",
            "| item_no | prompt | Taiwo | Jihwan |",
            "|---|---|---|---|",
        ]
        for _, i, prompt, x, y in rows:
            p = " ".join(prompt.replace("|", "/").split())
            lines.append(f"| {i} | {p[:80]}{'...' if len(p) > 80 else ''} | {x} | {y} |")
        lines.append("")
    (HERE / "agreement_report.md").write_text("\n".join(lines))

    for d, ag, ac1 in stats:
        print(f"{d:22s} {ag}/{n} ({100 * ag / n:.1f}%)  AC1={ac1:.3f}")
    print(f"overall {p_agree}/{tot} = {100 * p_agree / tot:.1f}%  pooled AC1={p_ac1:.3f}; disagreement rows={len(dis_rows)}")


if __name__ == "__main__":
    main()
