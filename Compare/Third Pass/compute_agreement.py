"""Final inter-rater agreement (Gwet's AC1 only, per professor directive
2026-09-11; Cohen's kappa intentionally not computed).

Sources, one per dimension:
  - deletion_location  -> Third Pass (items 51-100, revised codebook),
                          ./Taiwo_Coding.xlsx + ./Jihwan_Coding.xlsx
  - the other 6 dims   -> Second Pass (items 1-50),
                          ../Second Pass/{Taiwo,Jihwan}_Coding.xlsx
Why: deletion_location was the only dimension below AC1 0.80 in Second Pass;
its definition was tightened (see ../codebook.md, ../change.md) and re-tested
on a fresh 50 items. The other 6 dimensions were not re-coded.

Writes agreement_report.md and disagreements.csv here. For the 6 Second Pass
dimensions, adjudicated_value/adjudication_notes are carried over from
../Second Pass/disagreements.csv; deletion_location rows are left blank
(not yet adjudicated).

Usage: python3 compute_agreement.py
"""

import csv
from collections import Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
SECOND = HERE.parent / "Second Pass"
LOC = "deletion_location"
SECOND_PASS_DIMS = ["justification", "what", "mood", "tone", "verb", "accompanying_request"]
DIMENSIONS = [LOC, *SECOND_PASS_DIMS]
N = 50


def load(path: Path, dim: str) -> dict[int, tuple[str, str]]:
    ws = openpyxl.load_workbook(path)["Coding"]
    header = [c.value for c in ws[1]]
    col = header.index(dim)
    return {
        int(r[0]): (r[1], norm(r[col]))
        for r in ws.iter_rows(min_row=2, values_only=True)
        if r[0] is not None
    }


def norm(v) -> str:
    return "TRUE" if v is True else "FALSE" if v is False else str(v).strip().upper()


def gwet_ac1(a: list[str], b: list[str]) -> float:
    n, cats = len(a), sorted(set(a) | set(b))
    ca, cb = Counter(a), Counter(b)
    po = sum(x == y for x, y in zip(a, b)) / n
    pi = {c: (ca[c] + cb[c]) / (2 * n) for c in cats}
    pe = sum(p * (1 - p) for p in pi.values()) / (len(cats) - 1)
    return (po - pe) / (1 - pe)


def main() -> None:
    adjud = {
        (r["dimension"], r["item_no"]): r
        for r in csv.DictReader((SECOND / "disagreements.csv").open())
    }
    stats, dis_rows, pooled_t, pooled_j = [], [], [], []

    for dim in DIMENSIONS:
        folder = HERE if dim == LOC else SECOND
        t = load(folder / "Taiwo_Coding.xlsx", dim)
        j = load(folder / "Jihwan_Coding.xlsx", dim)
        items = sorted(t)
        assert t.keys() == j.keys() and len(items) == N, dim
        a, b = [t[i][1] for i in items], [j[i][1] for i in items]
        agree = sum(x == y for x, y in zip(a, b))
        stats.append((dim, agree, gwet_ac1(a, b), "Third" if dim == LOC else "Second",
                      f"{items[0]}-{items[-1]}"))
        pooled_t += [f"{dim}::{x}" for x in a]
        pooled_j += [f"{dim}::{x}" for x in b]
        for i in items:
            if t[i][1] != j[i][1]:
                adj = adjud.get((dim, str(i)), {}) if dim != LOC else {}
                dis_rows.append([dim, i, t[i][0], t[i][1], j[i][1],
                                 adj.get("adjudicated_value", ""), adj.get("adjudication_notes", "")])

    tot = len(pooled_t)
    p_agree = sum(x == y for x, y in zip(pooled_t, pooled_j))
    p_ac1 = gwet_ac1(pooled_t, pooled_j)

    with (HERE / "disagreements.csv").open("w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["dimension", "item_no", "prompt", "taiwo", "jihwan", "adjudicated_value", "adjudication_notes"])
        w.writerows(dis_rows)

    L = ["# Final Inter-Rater Agreement (Gwet's AC1, n=50 per dimension)", "",
         "| Dimension | Agreement | % | Gwet's AC1 |", "|---|---|---|---|"]
    for dim, ag, ac1, src, rng in stats:
        L.append(f"| {dim} | {ag}/{N} | {100*ag/N:.1f}% | {ac1:.3f} |")
    L += ["", f"**Overall micro agreement:** {p_agree}/{tot} = {100*p_agree/tot:.1f}% "
              f"(pooled Gwet's AC1 = {p_ac1:.3f})", "", "## Disagreements by dimension", ""]
    for dim in DIMENSIONS:
        rows = [r for r in dis_rows if r[0] == dim]
        L += [f"### {dim} ({len(rows)} disagreements)", "",
              "| item_no | prompt | Taiwo | Jihwan |", "|---|---|---|---|"]
        for r in rows:
            p = r[2].replace("|", "/")
            L.append(f"| {r[1]} | {p[:80]}{'...' if len(p) > 80 else ''} | {r[3].lower()} | {r[4].lower()} |")
        L.append("")
    (HERE / "agreement_report.md").write_text("\n".join(L))

    for s in stats:
        print(f"{s[0]:22s} {s[3]:6s} {s[1]}/{N}  AC1={s[2]:.3f}")
    print(f"overall {p_agree}/{tot} = {100*p_agree/tot:.1f}%  pooled AC1={p_ac1:.3f}; disagreements rows={len(dis_rows)}")


if __name__ == "__main__":
    main()
