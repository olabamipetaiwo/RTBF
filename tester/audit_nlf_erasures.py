"""One-time audit of every NL-forget cell already marked "erased" before
`assert_real_answer()` existed (see flows/base.py, added 2026-09-17).

The scheduler only ever checked that SOME reply appeared and stabilized,
never what it said -- so a rate-limit paywall, a generic service error,
or a blank/loading-only capture all got silently logged as a successful
erasure. Confirmed live: a Claude cell hit a hard session-limit paywall
and a Copilot cell got "I'm sorry, I'm having trouble responding to
requests right now" -- neither is a real answer, both were marked erased.

This script does NOT re-send anything. It revisits each cell's own
conversation (`injection_ref`, stored at injection time -- the same
conversation `_send_nl_forget` used) read-only, reads the reply already
sitting there, and classifies it with the exact same
`assert_real_answer()` logic the live pipeline now uses. Only genuine
non-answers get reverted (status -> injected, erased_at cleared, xlsx
RUN_STATUS -> INJECTED) so they're picked up fresh next run. A real
answer -- including an honest refusal ("I can't delete that") -- is left
exactly as-is; a refusal IS the result, not a failure.

Usage:
    python audit_nlf_erasures.py --dry-run   # classify only, no writes
    python audit_nlf_erasures.py             # classify AND revert bad ones
"""
from __future__ import annotations

import re
import sys
import time
import random
import traceback

import openpyxl

import config
import tracking
import run_cell as rc
from flows.claude import ClaudeFlow
from flows.chatgpt import ChatGPTFlow
from flows.gemini import GeminiFlow
from flows.copilot import CopilotFlow
from flows.perplexity import PerplexityFlow
from flows.deepseek import DeepSeekFlow

FLOW_REGISTRY = {
    "claude": ClaudeFlow,
    "chatgpt": ChatGPTFlow,
    "gemini": GeminiFlow,
    "copilot": CopilotFlow,
    "perplexity": PerplexityFlow,
    "deepseek": DeepSeekFlow,
}

# Method each flow uses to read the last assistant message -- not
# standardized across platforms (pre-existing naming, not touched here).
LAST_REPLY_METHOD = {
    "claude": "_last_reply_text",
    "chatgpt": "_last_assistant_text",
    "gemini": "_last_response_text",
    "copilot": "_last_reply_text",
    "perplexity": "_last_reply_text",
    "deepseek": "_last_reply_text",
}

NLF_COL_RUN_STATUS = 8  # see run_cell.py

LOG_PATH = config.TESTER_ROOT / "audit_nlf_erasures.log"


def log(msg: str) -> None:
    line = f"[{time.strftime('%Y-%m-%dT%H:%M:%S')}] {msg}"
    print(line, flush=True)
    with LOG_PATH.open("a") as f:
        f.write(line + "\n")


def already_confirmed_ok() -> set[str]:
    """Cell IDs this log already logged as 'OK' in an earlier run of this
    script -- skip re-visiting them on restart. Only 'OK' lines count;
    'BAD' cells were already reverted to injected so erased_nlf_cells()
    naturally excludes them regardless."""
    if not LOG_PATH.exists():
        return set()
    ok_ids = set()
    for line in LOG_PATH.read_text().splitlines():
        m = re.search(r"\] ([A-Z]+-NLF-I\d+) \([a-z]+\): OK --", line)
        if m:
            ok_ids.add(m.group(1))
    return ok_ids


def erased_nlf_cells() -> list[str]:
    data = tracking.load()
    skip = {p.strip() for p in (sys.argv[sys.argv.index("--skip") + 1].split(",") if "--skip" in sys.argv else [])}
    done = already_confirmed_ok()
    return sorted(
        cid for cid, e in data.items()
        if "-NLF-" in cid and e.get("status") == "erased"
        and rc.PLATFORM_PREFIX.get(cid[:2]) not in skip
        and cid not in done
    )


def revert_tracking(cell_id: str) -> None:
    data = tracking.load()
    entry = data.setdefault(cell_id, {})
    entry["status"] = "injected"
    entry["erased_at"] = None
    entry.pop("recall_due_at", None)
    tracking.save(data)


def revert_xlsx_row(cell_id: str, wb: openpyxl.Workbook) -> bool:
    ws = wb["NL FORGET"]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if row[0].value == cell_id:
            row[NLF_COL_RUN_STATUS - 1].value = "INJECTED"
            return True
    return False


def audit_one(cell_id: str, dry_run: bool) -> str:
    """Returns 'ok', 'reverted', or 'error' (error = couldn't audit, left untouched)."""
    platform = rc.PLATFORM_PREFIX.get(cell_id[:2])
    entry = tracking.get(cell_id)
    ref = entry.get("injection_ref") if entry else None
    if not ref:
        log(f"{cell_id}: SKIP -- no injection_ref stored, can't revisit")
        return "error"

    flow_cls = FLOW_REGISTRY[platform]
    method_name = LAST_REPLY_METHOD[platform]
    try:
        with flow_cls(session_label="nlforget") as flow:
            flow.page.goto(ref, wait_until="domcontentloaded")
            flow.page.wait_for_timeout(2500)
            reply = getattr(flow, method_name)()
        try:
            # Reuse the exact same check the live pipeline now applies --
            # needs a flow instance for `self.NON_ANSWER_SUBSTRINGS`, but
            # the check itself doesn't touch the page, safe to call after
            # __exit__ already closed the browser.
            flow.assert_real_answer(reply)
        except RuntimeError as e:
            log(f"{cell_id} ({platform}): BAD -- {e}")
            if not dry_run:
                revert_tracking(cell_id)
            return "reverted"
        log(f"{cell_id} ({platform}): OK -- {reply[:100]!r}")
        return "ok"
    except Exception as e:
        log(f"{cell_id} ({platform}): AUDIT ERROR (left untouched): {e}")
        log(traceback.format_exc())
        return "error"


def main() -> None:
    dry_run = "--dry-run" in sys.argv
    cells = erased_nlf_cells()
    log(f"=== audit start, {len(cells)} erased NLF cells to check, dry_run={dry_run} ===")

    results = {"ok": 0, "reverted": 0, "error": 0}
    reverted_ids: list[str] = []
    for i, cell_id in enumerate(cells):
        outcome = audit_one(cell_id, dry_run)
        results[outcome] += 1
        if outcome == "reverted":
            reverted_ids.append(cell_id)
        if i < len(cells) - 1:
            time.sleep(random.uniform(3, 8))

    log(f"=== audit done: {results['ok']} ok, {results['reverted']} reverted, {results['error']} errors ===")

    if reverted_ids and not dry_run:
        wb = openpyxl.load_workbook(config.MASTER_XLSX_PATH)
        for cid in reverted_ids:
            if not revert_xlsx_row(cid, wb):
                log(f"WARNING: {cid} reverted in tracking.json but not found in xlsx")
        wb.save(config.MASTER_XLSX_PATH)
        log(f"xlsx RUN_STATUS reverted for {len(reverted_ids)} cells")

    if reverted_ids:
        log("Reverted cell IDs: " + ", ".join(reverted_ids))


if __name__ == "__main__":
    main()
