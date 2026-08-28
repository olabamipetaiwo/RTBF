"""Generates the FILE substudy's real per-cell PDFs (DECISIONS Q14): one
fixed "Notes" template -- a few mundane filler lines plus the existing
per-cell disclosure line embedded among them -- rendered via Playwright's
page.pdf() (HTML -> PDF, no new dependency needed). Reuses each cell's
already-assigned referent+token (RTBF-Prompt/token_generator.py's
FILE_CONTENT_TEMPLATE, already sitting in the xlsx's "Disclosure
sentence" column for every FILE SUBSTUDY row) rather than regenerating
anything.

Usage:
    python generate_file_substudy_pdfs.py
"""

from __future__ import annotations

from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

import config

OUTPUT_DIR = config.TESTER_ROOT / "data" / "file_substudy_pdfs"

FILLER_BEFORE = [
    "Pick up dry cleaning before Friday",
]
FILLER_AFTER = [
    "Reply to book club group chat about next meeting time",
    "Look into new phone case options",
]


def build_notes_html(token_line: str) -> str:
    items = FILLER_BEFORE + [token_line] + FILLER_AFTER
    items_html = "".join(f"<li>{item}</li>" for item in items)
    return f"""<!doctype html>
<html><head><meta charset="utf-8"><style>
body {{ font-family: Helvetica, Arial, sans-serif; padding: 40px; }}
h1 {{ font-size: 20px; }}
li {{ font-size: 14px; line-height: 1.8; }}
</style></head>
<body><h1>Notes</h1><ul>{items_html}</ul></body></html>"""


def load_file_substudy_cells() -> list[tuple[str, str]]:
    """Returns [(cell_id, disclosure_line), ...] for every FILE SUBSTUDY row."""
    wb = openpyxl.load_workbook(config.MASTER_XLSX_PATH, data_only=True)
    ws = wb["FILE SUBSTUDY"]
    cells = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        cell_id = row[0]
        disclosure_line = row[12]
        if cell_id and disclosure_line:
            cells.append((cell_id, disclosure_line))
    return cells


def main() -> None:
    cells = load_file_substudy_cells()
    print(f"Generating {len(cells)} PDFs into {OUTPUT_DIR}")
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    with sync_playwright() as p:
        browser = p.chromium.launch()
        page = browser.new_page()
        for cell_id, disclosure_line in cells:
            page.set_content(build_notes_html(disclosure_line))
            out_path = OUTPUT_DIR / f"{cell_id}.pdf"
            page.pdf(path=str(out_path), format="Letter")
            print(f"  {cell_id} -> {out_path.name}")
        browser.close()

    print("Done.")


if __name__ == "__main__":
    main()
